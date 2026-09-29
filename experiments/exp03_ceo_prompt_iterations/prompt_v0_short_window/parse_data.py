import os
import json
import pandas as pd
from docx import Document
import re

def parse_gold():
    xls = pd.ExcelFile('gold.xlsx')
    gold_data = []
    
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet, header=1)
        for _, row in df.iterrows():
            item = row.to_dict()
            clean_item = {str(k): (v if pd.notna(v) else None) for k, v in item.items()}
            gold_data.append(clean_item)
            
    with open('Gold_Parsed/gold.json', 'w', encoding='utf-8') as f:
        json.dump(gold_data, f, indent=4, ensure_ascii=False)
    print(f"Parsed {len(gold_data)} gold items.")

def parse_docx(filepath):
    doc = Document(filepath)
    items = []
    
    # 1. Parse Tables
    for table in doc.tables:
        if len(table.rows) > 1:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            # Check if it looks like our table
            if any(h in headers for h in ['Hospital', 'Issue', 'Date']):
                for row in table.rows[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    if len(cells) == len(headers):
                        item = dict(zip(headers, cells))
                        items.append(item)
    
    # 2. Parse Paragraphs (ChatGPT block style)
    current_item = {}
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        if re.match(r'^\d+$', text) or re.match(r'^\d+\.$', text):
            if current_item and 'Hospital' in current_item:
                items.append(current_item)
            current_item = {}
            continue
            
        # ChatGPT / Claude format
        if '→' in text and len(text.split('→')) == 2:
            parts = text.split('→', 1)
            key = parts[0].strip()
            val = parts[1].strip()
            current_item[key] = val
        # Perplexity inline format
        elif '→' in text and len(text.split('→')) > 2:
            parts = [p.strip() for p in text.split('→')]
            if len(parts) >= 6: # Hospital -> Issue -> Blurb -> Date -> Source -> Source link
                items.append({
                    'Hospital': parts[0],
                    'Issue': parts[1],
                    'Paragraph blurb': parts[2],
                    'Date': parts[3],
                    'Source': parts[4],
                    'Source link': parts[5] if len(parts) > 5 else None
                })
        else:
            # Maybe it's a scope note or intro text, ignore
            pass
            
    if current_item and 'Hospital' in current_item:
        items.append(current_item)
        
    return items

def parse_all_models():
    base_dir = '../raw_responses'
    models = [d for d in os.listdir(base_dir) if os.path.isdir(os.path.join(base_dir, d))]
    all_responses = {}
    
    for model in models:
        model_dir = os.path.join(base_dir, model)
        all_responses[model] = []
        for root, dirs, files in os.walk(model_dir):
            for file in files:
                if file.endswith('.docx') and not file.startswith('~'):
                    filepath = os.path.join(root, file)
                    run_items = parse_docx(filepath)
                    for item in run_items:
                        item['_meta'] = {
                            'file': file,
                            'path': filepath,
                            'run': file.replace('.docx', '')
                        }
                    all_responses[model].extend(run_items)
                    
    for model, items in all_responses.items():
        out_path = os.path.join('Parsed_Responses', f"{model}.json")
        with open(out_path, 'w', encoding='utf-8') as f:
            json.dump(items, f, indent=4, ensure_ascii=False)
        print(f"Parsed {len(items)} items for {model}.")

if __name__ == '__main__':
    parse_gold()
    parse_all_models()
