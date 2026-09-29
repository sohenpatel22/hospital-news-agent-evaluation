from docx import Document
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = "../raw_responses/MS Copilot Agent/Mackenzie Health/Run 1.docx"
doc = Document(path)
print(f"Number of tables: {len(doc.tables)}")
if len(doc.tables) > 0:
    for row in doc.tables[0].rows[:3]:
        print([cell.text for cell in row.cells])
