import asyncio
import os
import json
from docx import Document
import pandas as pd
from google.antigravity import Agent, LocalAgentConfig

async def extract_gold_items():
    xls = pd.ExcelFile('gold.xlsx')
    gold_data = []
    
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet, header=1)
        for _, row in df.iterrows():
            item = row.to_dict()
            clean_item = {str(k): (v if pd.notna(v) else None) for k, v in item.items()}
            gold_data.append(clean_item)
            
    # Since gold is already tabulated, we don't necessarily need the LLM to parse it from unstructured text. 
    # But to follow instructions "parsing gold dataset and again use LLM to match both", 
    # the matching part is strictly LLM. The gold dataset is already in Excel.
    with open('Gold_Parsed/gold.json', 'w', encoding='utf-8') as f:
        json.dump(gold_data, f, indent=4, ensure_ascii=False)
    print(f"Parsed {len(gold_data)} gold items deterministically as it is an Excel file.")

async def extract_items_from_text(text, agent):
    prompt = f"""
Extract all hospital performance management events (appointments, departures, governance changes) from the text below.
Return ONLY a valid JSON array of objects, with each object having exactly these keys:
- "Hospital" (string)
- "Issue" (string)
- "Paragraph blurb" (string)
- "Date" (string)
- "Source" (string)
- "Source link" (string)

If any field is missing, use null.
Do not output anything else but the JSON array.

Text:
{text}
"""
    try:
        response = await agent.chat(prompt)
        content = response.text
        # extract JSON block
        if "```json" in content:
            content = content.split("```json")[1].split("```")[0].strip()
        elif "```" in content:
            content = content.split("```")[1].strip()
            
        items = json.loads(content)
        if isinstance(items, list):
            return items
        else:
            return []
    except Exception as e:
        print(f"Error parsing with LLM: {e}")
        return []

async def parse_all_models():
    config = LocalAgentConfig(system_instructions="You are an expert data extractor. Extract exactly the requested JSON.")
    base_dir = '../raw_responses'
    models = [d for d in os.listdir(base_dir) if os.path.isdir(os.path.join(base_dir, d))]
    
    all_responses = {}
    
    async with Agent(config) as agent:
        for model in models:
            model_dir = os.path.join(base_dir, model)
            all_responses[model] = []
            print(f"Processing model: {model}")
            
            for root, dirs, files in os.walk(model_dir):
                for file in files:
                    if file.endswith('.docx') and not file.startswith('~'):
                        # Process only Run 1.docx files (or all runs for Copilot for consistency)
                        if file != 'Run 1.docx' and not (model == 'MS Copilot Agent' and file.startswith('Run ')):
                            continue
                            
                        filepath = os.path.join(root, file)
                        doc = Document(filepath)
                        
                        text = []
                        for para in doc.paragraphs:
                            if para.text.strip(): text.append(para.text.strip())
                        for table in doc.tables:
                            for row in table.rows:
                                text.append(" | ".join([cell.text.strip() for cell in row.cells]))
                                
                        full_text = "\n".join(text)
                        if not full_text: continue
                        
                        run_items = await extract_items_from_text(full_text, agent)
                        
                        for item in run_items:
                            item['_meta'] = {
                                'file': file,
                                'path': filepath,
                                'run': file.replace('.docx', '')
                            }
                        all_responses[model].extend(run_items)
                        print(f"  - Parsed {len(run_items)} items from {filepath}")
                        
            out_path = os.path.join('Parsed_Responses', f"{model}_llm.json")
            with open(out_path, 'w', encoding='utf-8') as f:
                json.dump(all_responses[model], f, indent=4, ensure_ascii=False)
            print(f"Saved {len(all_responses[model])} items for {model}.")

async def main():
    await extract_gold_items()
    await parse_all_models()

if __name__ == '__main__':
    asyncio.run(main())
