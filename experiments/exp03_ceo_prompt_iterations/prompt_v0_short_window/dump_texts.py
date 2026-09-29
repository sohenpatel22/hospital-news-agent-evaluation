import os
from docx import Document

base_dir = '../raw_responses'
models = ['ChatGPT', 'Claude', 'Gemini', 'MS Copilot Agent', 'Perplexity']

with open('all_texts.txt', 'w', encoding='utf-8') as f_out:
    for model in models:
        model_dir = os.path.join(base_dir, model)
        if not os.path.isdir(model_dir): continue
        for root, dirs, files in os.walk(model_dir):
            for file in files:
                if file == 'Run 1.docx':
                    filepath = os.path.join(root, file)
                    doc = Document(filepath)
                    text = []
                    for para in doc.paragraphs:
                        if para.text.strip(): text.append(para.text.strip())
                    for table in doc.tables:
                        for row in table.rows:
                            text.append(" | ".join([c.text.strip() for c in row.cells]))
                    f_out.write(f"\n\n--- MODEL: {model} | FILE: {filepath} ---\n")
                    f_out.write("\n".join(text))
