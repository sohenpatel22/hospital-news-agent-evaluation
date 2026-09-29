"""
06_final_report_llm.py
======================
Compiles all LLM evaluation results into a single formatted Excel workbook.
  Sheet 1: Overall Summary (all metrics, all models)
  Sheet 2: Task Summary (per task breakdown)
  Sheet 3: Item-Level (all hospital rows with all metrics)
  Sheet 4: LLM Judge Details (reasoning text)
  Sheet 5: Citation Analysis
  Sheet 6: Metric Definitions
"""

import json
import pathlib
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = pathlib.Path(__file__).parent
METRICS_DIR = BASE / "metrics"
JUDGE_DIR = BASE / "judge_scores"
VIZ_DIR = BASE / "visualizations"
OUT_DIR = BASE / "final"

METRIC_DEFINITIONS = [
    ("Recall",           "Proportion of gold-standard events correctly identified by the model."),
    ("Precision",        "Proportion of model-returned events that are correct (not hallucinated)."),
    ("F1 Score",         "Harmonic mean of Recall and Precision. Balances both."),
    ("Relevance",        "LLM-judged (0-3 → 0-1): Does the response directly answer the user's intent?"),
    ("Groundedness",     "LLM-judged (0-3 → 0-1): Are factual claims supported by cited sources?"),
    ("Completeness",     "LLM-judged (0-1): What fraction of gold items does the response cover?"),
    ("Answer Relevancy", "RAGAS-inspired (0-1): How focused/on-topic is the response?"),
    ("Citation Accuracy","HTTP-validated (0-1): Proportion of cited URLs that resolve successfully."),
    ("Abstention Correct","Deterministic: Did model correctly say 'none found' for zero-event hospitals?"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")  # dark blue
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
COPILOT_FILL = PatternFill("solid", fgColor="E8F0FE")  # light blue highlight
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
NORMAL_FONT = Font(size=9)
BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC")
)

def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BORDER

def auto_width(ws, max_width=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 8), max_width)

def run_report():
    print("--- Stage 6: Final Excel Report ---")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_file = OUT_DIR / "Evaluation_Results_LLM.xlsx"
    
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        
        # --- Sheet 1: Overall Summary ---
        overall_path = METRICS_DIR / "overall_summary_llm.csv"
        if overall_path.exists():
            df_overall = pd.read_csv(overall_path, index_col=0)
            df_overall.reset_index(inplace=True)
            df_overall.rename(columns={"index": "Model"}, errors="ignore")
            df_overall.to_excel(writer, sheet_name="Overall Summary", index=True)
            ws = writer.sheets["Overall Summary"]
            style_header(ws)
            ws.freeze_panes = "B2"
            auto_width(ws)
        
        # --- Sheet 2: Task Summary ---
        task_path = METRICS_DIR / "tool_summary_llm.csv"
        if task_path.exists():
            df_task = pd.read_csv(task_path)
            df_task.to_excel(writer, sheet_name="Task Summary", index=False)
            ws = writer.sheets["Task Summary"]
            style_header(ws)
            ws.freeze_panes = "A2"
            auto_width(ws)
            # Highlight Copilot rows
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                if row[1].value == "MS Copilot Agent":
                    for cell in row:
                        cell.fill = COPILOT_FILL
        
        # --- Sheet 3: Item Level (Combined Scores) ---
        combined_path = METRICS_DIR / "combined_scores.csv"
        if combined_path.exists():
            df_combined = pd.read_csv(combined_path)
            # Drop very long text columns for readability
            drop_cols = ["relevance_reason","groundedness_reason","completeness_reason",
                         "answer_relevancy_reason","generated_questions"]
            df_combined = df_combined.drop(columns=[c for c in drop_cols if c in df_combined.columns])
            df_combined.to_excel(writer, sheet_name="Item Level", index=False)
            ws = writer.sheets["Item Level"]
            style_header(ws)
            ws.freeze_panes = "D2"
            auto_width(ws)
        
        # --- Sheet 4: LLM Judge Details ---
        judge_rows = []
        if JUDGE_DIR.exists():
            for task_dir in JUDGE_DIR.iterdir():
                if not task_dir.is_dir(): continue
                for model_dir in task_dir.iterdir():
                    if not model_dir.is_dir(): continue
                    for f in model_dir.glob("*.json"):
                        try:
                            data = json.loads(f.read_text(encoding="utf-8"))
                            judge_rows.append({
                                "task": data.get("task"),
                                "model": data.get("model"),
                                "hospital": data.get("hospital"),
                                "relevance": data.get("relevance"),
                                "relevance_reason": data.get("relevance_reason",""),
                                "groundedness": data.get("groundedness"),
                                "groundedness_reason": data.get("groundedness_reason",""),
                                "completeness": data.get("completeness"),
                                "completeness_reason": data.get("completeness_reason",""),
                                "missed_items": "; ".join(data.get("missed_items",[])),
                                "answer_relevancy": data.get("answer_relevancy"),
                                "answer_relevancy_reason": data.get("answer_relevancy_reason",""),
                            })
                        except Exception:
                            pass
        if judge_rows:
            df_judge = pd.DataFrame(judge_rows)
            df_judge.to_excel(writer, sheet_name="LLM Judge Details", index=False)
            ws = writer.sheets["LLM Judge Details"]
            style_header(ws)
            ws.freeze_panes = "D2"
            auto_width(ws, max_width=80)
        
        # --- Sheet 5: Citation Analysis ---
        cite_path = METRICS_DIR / "citation_analysis.csv"
        if cite_path.exists():
            df_cite = pd.read_csv(cite_path)
            df_cite.to_excel(writer, sheet_name="Citation Analysis", index=False)
            ws = writer.sheets["Citation Analysis"]
            style_header(ws)
            ws.freeze_panes = "A2"
            auto_width(ws)
        
        # --- Sheet 6: Metric Definitions ---
        df_defs = pd.DataFrame(METRIC_DEFINITIONS, columns=["Metric", "Definition"])
        df_defs.to_excel(writer, sheet_name="Metric Definitions", index=False)
        ws = writer.sheets["Metric Definitions"]
        style_header(ws)
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 80
    
    print(f"Done. Report saved to {out_file}")

if __name__ == "__main__":
    run_report()
