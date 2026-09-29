import os
import csv
from docx import Document
import re

def extract_docx_to_csv():
    print("--- Phase 3: Document Ingestion ---")
    base_dir = r"../raw_responses/Prompt V0"
    out_file = "raw_outputs/raw_outputs.csv"
    
    rows = []
    
    for task_folder in os.listdir(base_dir):
        task_path = os.path.join(base_dir, task_folder)
        if not os.path.isdir(task_path): continue
        
        # normalize task name
        task = task_folder.lower()
        if "phipa" in task: task = "phipa"
        elif "ceo" in task: task = "ceo"
        elif "supervisor" in task: task = "supervisors"
        
        for model_folder in os.listdir(task_path):
            model_path = os.path.join(task_path, model_folder)
            if not os.path.isdir(model_path): continue
            
            for hospital_folder in os.listdir(model_path):
                hospital_path = os.path.join(model_path, hospital_folder)
                if not os.path.isdir(hospital_path): continue
                
                for file_name in os.listdir(hospital_path):
                    if not file_name.endswith('.docx') or file_name.startswith('~'): continue
                    
                    file_path = os.path.join(hospital_path, file_name)
                    
                    # Extract run number (e.g. "Run 1.docx" -> 1)
                    m = re.search(r'Run\s*(\d+)', file_name, re.IGNORECASE)
                    run_num = int(m.group(1)) if m else 1
                    
                    # For commercial LLMs, use Run 1 only
                    is_copilot = ("copilot" in model_folder.lower())
                    if not is_copilot and run_num != 1:
                        continue
                        
                    # Read docx
                    try:
                        doc = Document(file_path)
                    except Exception as e:
                        print(f"Error reading {file_path}: {e}")
                        continue
                        
                    # Extract paragraphs
                    text_parts = []
                    for p in doc.paragraphs:
                        text = p.text.strip()
                        if text:
                            text_parts.append(text)
                            
                    # Extract tables as markdown-like text
                    for table in doc.tables:
                        for row in table.rows:
                            row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                            text_parts.append(" | ".join(row_data))
                            
                    full_text = "\n".join(text_parts)
                    
                    rows.append({
                        'task': task,
                        'model': model_folder,
                        'hospital': hospital_folder,
                        'run_number': run_num,
                        'raw_text': full_text,
                        'file_path': file_path
                    })
                    
    # Write to CSV
    with open(out_file, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=['task', 'model', 'hospital', 'run_number', 'raw_text', 'file_path'])
        writer.writeheader()
        writer.writerows(rows)
        
    print(f"Extracted {len(rows)} documents.")
    if len(rows) == 162:
        print("PASS: Document count is exactly 162.")
    else:
        print(f"WARNING: Expected 162 documents, got {len(rows)}.")
        
    print(f"Saved to {out_file}")

if __name__ == "__main__":
    extract_docx_to_csv()
